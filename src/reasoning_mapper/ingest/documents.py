# SPDX-License-Identifier: Apache-2.0
# Copyright (C) 2026 Tucuxi, Inc.
"""Documents the expert wrote: memos, playbooks, training decks, post-mortems.

Decks deserve a warning. Presenters often put the worked answer in slide bodies
and the reasoning in what they SAID, which no file records. Slide text and notes
are both extracted, but a deck is usually the answer without the route - pair it
with a think-aloud session over the same material (docs/THINK_ALOUD.md).
"""

from __future__ import annotations

from pathlib import Path

from . import insert_text

TEXT_EXT = {".txt", ".md", ".markdown", ".rst"}


def read_docx(p: Path) -> str:
    import docx
    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(x for x in parts if x.strip())


def read_pptx(p: Path) -> str:
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(str(p)).slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes
                 if sh.has_text_frame and sh.text_frame.text.strip()]
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        out.append(f"[slide {i}]\n" + "\n".join(texts)
                   + (f"\n[speaker notes]\n{notes}" if notes else ""))
    return "\n\n".join(out)


def read_pdf(p: Path) -> str:
    from pypdf import PdfReader
    return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)


READERS = {".docx": read_docx, ".pptx": read_pptx, ".pdf": read_pdf}


def read_any(p: Path) -> str | None:
    ext = p.suffix.lower()
    if ext in TEXT_EXT:
        return p.read_text(errors="replace")
    if ext in READERS:
        try:
            return READERS[ext](p)
        except ImportError as e:
            raise SystemExit(
                f"{ext} support needs extras: pip install 'reasoning-mapper[docs]'") from e
    return None


def ingest_documents(conn, source: str, path: Path, authored_by_expert: bool) -> str:
    files = [path] if path.is_file() else sorted(p for p in path.rglob("*") if p.is_file())
    new = skipped = failed = 0
    for p in files:
        try:
            text = read_any(p)
        except SystemExit:
            raise
        except Exception:
            failed += 1
            continue
        if not text or not text.strip():
            skipped += 1
            continue
        if insert_text(conn, "documents", source, p, p.stem, text,
                       authored_by_expert=int(authored_by_expert)):
            new += 1
    return f"{source}: {new} new documents, {skipped} empty/unsupported, {failed} unreadable"
